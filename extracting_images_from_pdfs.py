import os
import io
import fitz  # PyMuPDF
import cv2
import pytesseract
import numpy as np
from PIL import Image
from pptx import Presentation

from img2table.document import PDF as Img2TablePDF
from img2table.ocr import TesseractOCR
# import pytesseract

pytesseract.pytesseract.tesseract_cmd = (
    r"C:\Program Files\Tesseract-OCR\tesseract.exe"
)


# ---------------- CONFIG ----------------
PDF_ROOT = r"C:\\Users\\smwaj\\fyp_text\\text_pdfs"
OUTPUT_ROOT = r"C:\\Users\\smwaj\\fyp_text\\images"
LOGO_TEMPLATES_FOLDER = r"C:\\Users\\smwaj\\fyp_text\\text_pdfs\\logo_templates"

TEXT_LENGTH_THRESHOLD = 10
MIN_IMAGE_SIZE = 100
TEMPLATE_MATCH_THRESHOLD = 0.95
MAX_ASPECT_RATIO = 3.0
LOW_DETAIL_STD_THRESHOLD = 20

OCR_LANGS = "eng+urd"

# ---------------- LOGO UTILITIES ----------------
def load_logo_templates(folder_path):
    templates = []
    if not os.path.exists(folder_path):
        return templates
    for file in os.listdir(folder_path):
        if file.lower().endswith((".png", ".jpg", ".jpeg")):
            templates.append(Image.open(os.path.join(folder_path, file)).convert("RGB"))
    return templates

def is_similar_to_logo_template(pil_img, template_img):
    img_gray = cv2.cvtColor(np.array(pil_img), cv2.COLOR_RGB2GRAY)
    tpl_gray = cv2.cvtColor(np.array(template_img), cv2.COLOR_RGB2GRAY)
    ih, iw = img_gray.shape
    th, tw = tpl_gray.shape
    if th > ih or tw > iw:
        return False
    for scale in [1.0, 0.9, 0.8, 0.7]:
        resized = cv2.resize(tpl_gray, (int(tw * scale), int(th * scale)))
        res = cv2.matchTemplate(img_gray, resized, cv2.TM_CCOEFF_NORMED)
        if res.max() >= TEMPLATE_MATCH_THRESHOLD:
            return True
    return False

def is_similar_to_any_logo_template(pil_img, templates):
    return any(is_similar_to_logo_template(pil_img, tpl) for tpl in templates)

# ---------------- IMAGE FILTERS ----------------
def is_low_detail_image(pil_img):
    gray = cv2.cvtColor(np.array(pil_img), cv2.COLOR_RGB2GRAY)
    return gray.std() < LOW_DETAIL_STD_THRESHOLD

def image_contains_text(pil_img):
    gray = cv2.cvtColor(np.array(pil_img), cv2.COLOR_RGB2GRAY)
    text = pytesseract.image_to_string(gray, lang=OCR_LANGS)
    return len(text.strip()) >= TEXT_LENGTH_THRESHOLD

def page_contains_significant_text(pil_img):
    gray = cv2.cvtColor(np.array(pil_img), cv2.COLOR_RGB2GRAY)
    text = pytesseract.image_to_string(gray, lang=OCR_LANGS)
    return len(text.strip()) > 100

# ---------------- TABLE EXTRACTION ----------------
def extract_tables_from_pdf(pdf_path, output_folder, pdf_name):
    os.makedirs(output_folder, exist_ok=True)
    ocr = TesseractOCR(lang=OCR_LANGS)
    pdf = Img2TablePDF(src=pdf_path)
    table_pages = set()

    try:
        extracted_tables = pdf.extract_tables(
            ocr=ocr,
            implicit_rows=True,
            borderless_tables=True,
            min_confidence=50
        )
    except Exception as e:
        print(f"⚠️ Error extracting tables from {pdf_name}: {e}")
        return table_pages

    for page_number, tables in extracted_tables.items():
        if not tables:
            continue
        table_pages.add(page_number)
        page_img = pdf.images[page_number]

        for idx, table in enumerate(tables, start=1):
            x1, y1, x2, y2 = table.bbox.x1, table.bbox.y1, table.bbox.x2, table.bbox.y2
            table_img = page_img[y1:y2, x1:x2]
            if table_img.size == 0:
                continue
            cv2.imwrite(os.path.join(output_folder, f"{pdf_name}_table_page{page_number+1}_{idx}.png"), table_img)

    return table_pages

# ---------------- PDF IMAGE EXTRACTION ----------------
def extract_images_from_pdf(pdf_path, output_folder, logo_templates):
    pdf_name = os.path.splitext(os.path.basename(pdf_path))[0]
    table_pages = extract_tables_from_pdf(pdf_path, output_folder, pdf_name)
    doc = fitz.open(pdf_path)
    img_index = 0

    for page_number, page in enumerate(doc, start=1):
        # Embedded images
        for img in page.get_images(full=True):
            xref = img[0]
            base_img = doc.extract_image(xref)
            img_bytes = base_img["image"]
            img_ext = base_img["ext"].lower()
            if img_ext not in ["jpg", "jpeg", "png"]:
                continue
            pil_img = Image.open(io.BytesIO(img_bytes)).convert("RGB")
            w, h = pil_img.size
            ar = w / h
            if w < MIN_IMAGE_SIZE or h < MIN_IMAGE_SIZE or ar > MAX_ASPECT_RATIO:
                continue
            if is_low_detail_image(pil_img):
                continue
            if is_similar_to_any_logo_template(pil_img, logo_templates):
                continue
            if image_contains_text(pil_img):
                img_index += 1
                pil_img.save(os.path.join(output_folder, f"{pdf_name}_img_page{page_number}_{img_index}.{img_ext}"))

        # Full-page diagrams
        if (page_number - 1) in table_pages:
            continue
        mat = fitz.Matrix(2.0, 2.0)
        pix = page.get_pixmap(matrix=mat)
        page_img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        if is_low_detail_image(page_img) or is_similar_to_any_logo_template(page_img, logo_templates) or page_contains_significant_text(page_img):
            continue
        page_img.save(os.path.join(output_folder, f"{pdf_name}_diagram_page{page_number}.png"))

    doc.close()

# ---------------- PPTX IMAGE EXTRACTION ----------------
def extract_images_from_pptx(pptx_path, output_folder, logo_templates):
    pptx_name = os.path.splitext(os.path.basename(pptx_path))[0]
    prs = Presentation(pptx_path)
    img_index = 0

    for slide_idx, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.shape_type == 13:  # picture
                img_bytes = shape.image.blob
                img_ext = shape.image.ext.lower()
                if img_ext not in ["jpg", "jpeg", "png"]:
                    continue
                pil_img = Image.open(io.BytesIO(img_bytes)).convert("RGB")
                w, h = pil_img.size
                ar = w / h
                if w < MIN_IMAGE_SIZE or h < MIN_IMAGE_SIZE or ar > MAX_ASPECT_RATIO:
                    continue
                if is_low_detail_image(pil_img):
                    continue
                if is_similar_to_any_logo_template(pil_img, logo_templates):
                    continue
                if image_contains_text(pil_img):
                    img_index += 1
                    pil_img.save(os.path.join(output_folder, f"{pptx_name}_img_slide{slide_idx}_{img_index}.{img_ext}"))

# ---------------- MAIN ----------------
def main():
    logo_templates = load_logo_templates(LOGO_TEMPLATES_FOLDER)

    for root, _, files in os.walk(PDF_ROOT):
        for file in files:
            file_path = os.path.join(root, file)
            rel_path = os.path.relpath(root, PDF_ROOT)
            output_folder = os.path.join(OUTPUT_ROOT, rel_path)
            os.makedirs(output_folder, exist_ok=True)

            base_name = os.path.splitext(file)[0]
            # Check if images already exist
            existing_images = [f for f in os.listdir(output_folder) if f.startswith(base_name)]
            if existing_images:
                print(f"ℹ️ Skipping {file}, images already extracted.")
                continue

            if file.lower().endswith(".pdf"):
                extract_images_from_pdf(file_path, output_folder, logo_templates)
            elif file.lower().endswith(".pptx"):
                extract_images_from_pptx(file_path, output_folder, logo_templates)

    print("✅ Extraction completed successfully.")

if __name__ == "__main__":
    main()